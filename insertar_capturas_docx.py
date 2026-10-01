from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


BASE = Path(r"C:\Users\ander\source\repos\Laboratorio-07")
ORIGINAL = Path(r"C:\Users\ander\Downloads\GLAB-S07-EAREVALO-2026-2.docx")
OUTPUT = BASE / "GLAB-S07-EAREVALO-2026-2-completado.docx"
CAPTURES = BASE / "capturas"

FIGURES = [
    (
        "Vista de mantenimiento de libros",
        CAPTURES / "01-libros.png",
        "Figura 1. Consulta y mantenimiento de libros registrados en la biblioteca.",
    ),
    (
        "Vista de mantenimiento de socios",
        CAPTURES / "02-socios.png",
        "Figura 2. Consulta y mantenimiento de socios de la biblioteca.",
    ),
    (
        "Vista de préstamos y devoluciones",
        CAPTURES / "03-prestamos-devoluciones.png",
        "Figura 3. Registro de préstamos y gestión de devoluciones pendientes.",
    ),
    (
        "Vista del reporte de préstamos",
        CAPTURES / "04-reporte.png",
        "Figura 4. Reporte de préstamos filtrado por un rango de fechas.",
    ),
]


def configure_paragraph_spacing(paragraph, before=0, after=0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def main():
    if not ORIGINAL.exists():
        raise FileNotFoundError(f"No se encontró el documento original: {ORIGINAL}")
    for _, image_path, _ in FIGURES:
        if not image_path.exists():
            raise FileNotFoundError(f"No se encontró la captura: {image_path}")

    document = Document(ORIGINAL)

    document.add_page_break()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure_paragraph_spacing(title, after=8)
    run = title.add_run("EVIDENCIAS DE LA APLICACIÓN")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(31, 78, 121)

    intro = document.add_paragraph(
        "Las siguientes capturas evidencian el funcionamiento de las vistas principales de la aplicación de gestión de biblioteca."
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    configure_paragraph_spacing(intro, after=12, line=1.15)
    for intro_run in intro.runs:
        intro_run.font.name = "Arial"
        intro_run.font.size = Pt(10)

    section = document.sections[-1]
    available_width = section.page_width - section.left_margin - section.right_margin
    max_width = min(available_width, Inches(6.7))

    for index, (heading_text, image_path, caption_text) in enumerate(FIGURES):
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.page_break_before = bool(index)
        heading.paragraph_format.left_indent = Inches(0)
        heading.paragraph_format.right_indent = Inches(0)
        heading.paragraph_format.first_line_indent = Inches(0)
        configure_paragraph_spacing(heading, after=8)
        heading_run = heading.add_run(heading_text)
        heading_run.bold = True
        heading_run.font.name = "Arial"
        heading_run.font.size = Pt(13)
        heading_run.font.color.rgb = RGBColor(31, 78, 121)

        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.keep_with_next = True
        configure_paragraph_spacing(image_paragraph, after=5)
        image_paragraph.add_run().add_picture(str(image_path), width=max_width)

        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        configure_paragraph_spacing(caption, after=6)
        caption_run = caption.add_run(caption_text)
        caption_run.italic = True
        caption_run.font.name = "Arial"
        caption_run.font.size = Pt(9)
        caption_run.font.color.rgb = RGBColor(89, 89, 89)

    document.core_properties.title = "Laboratorio 07 - Evidencias de la aplicación Biblioteca"
    document.core_properties.subject = "Capturas de las vistas de la aplicación"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
